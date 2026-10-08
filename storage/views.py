from django.shortcuts import render, redirect
from django.contrib.auth.models import User
from django.contrib.auth import authenticate, login as auth_login
from django.contrib.auth import logout
from django.contrib.auth.decorators import login_required
import hashlib
import mammoth
from pptx import Presentation
from .models import UploadedFile


def home(request):
    return render(request, 'home.html')


def login_view(request):

    if request.method == 'POST':

        username = request.POST.get('username')
        password = request.POST.get('password')

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None:
            auth_login(request, user)
            return redirect('dashboard')

        else:
            return render(request, 'login.html', {
                'message': 'Invalid username or password!'
            })

    return render(request, 'login.html')
def register_view(request):

    if request.method == 'POST':

        username = request.POST.get('username')
        email = request.POST.get('email')
        password = request.POST.get('password')
        confirm_password = request.POST.get('confirm_password')

        if password != confirm_password:

            return render(request, 'register.html', {
                'message': 'Passwords do not match!'
            })

        if User.objects.filter(username=username).exists():

            return render(request, 'register.html', {
                'message': 'Username already exists! Please choose another username.'
            })

        User.objects.create_user(
            username=username,
            email=email,
            password=password
        )

        return render(request, 'register.html', {
            'message': 'Registration successful!'
        })

    return render(request, 'register.html')

@login_required
def dashboard(request):

    if request.method == 'POST':

        uploaded_file = request.FILES.get('file')

        if uploaded_file:

            # Read uploaded file content
            uploaded_content = uploaded_file.read()

            # Reset file pointer
            uploaded_file.seek(0)

            # Check every existing file of this user
            existing_files = UploadedFile.objects.filter(
                user=request.user
            )

            duplicate_found = False

            for existing_file in existing_files:

                try:
                    existing_file.file.open('rb')
                    existing_content = existing_file.file.read()
                    existing_file.file.close()
                    print("UPLOADED:", len(uploaded_content))
                    print("EXISTING:", len(existing_content))
                    print("SAME:", uploaded_content == existing_content)
                    # Compare exact content
                    if uploaded_content == existing_content:
                        duplicate_found = True
                        break

                except (FileNotFoundError, ValueError):
                    continue

            # If duplicate, do not save
            if duplicate_found:

                files = UploadedFile.objects.filter(
                    user=request.user
                ).order_by('-uploaded_at')

                total_files = files.count()

                storage_used = round(
                    sum(file.file_size for file in files),
                    2
                )

                total_carbon = round(
                    sum(file.carbon_impact for file in files),
                    4
                )

                if total_carbon < 1:
                    carbon_value = round(
                        total_carbon * 1000,
                        2
                    )
                    carbon_unit = 'mg CO₂'
                else:
                    carbon_value = round(
                        total_carbon,
                        2
                    )
                    carbon_unit = 'g CO₂'

                return render(
                    request,
                    'dashboard.html',
                    {
                        'files': files,
                        'total_files': total_files,
                        'storage_used': storage_used,
                        'total_carbon': carbon_value,
                        'carbon_unit': carbon_unit,
                        'message': 'Duplicate file detected! This file was not stored.'
                    }
                )

            # File size
            file_size = uploaded_file.size

            # Convert bytes to MB
            file_size_mb = round(
                file_size / (1024 * 1024),
                2
            )

            # Convert MB to GB
            file_size_gb = file_size_mb / 1024

            # Carbon estimation
            STORAGE_ENERGY_FACTOR = 0.06
            GRID_CARBON_FACTOR = 442

            carbon_impact = round(
                file_size_gb
                * STORAGE_ENERGY_FACTOR
                * GRID_CARBON_FACTOR,
                4
            )

            # Store file
            UploadedFile.objects.create(
                user=request.user,
                file=uploaded_file,
                file_name=uploaded_file.name,
                file_size=file_size_mb,
                carbon_impact=carbon_impact,
                file_hash=''
            )

            return redirect('dashboard')

    # Get user's files
    files = UploadedFile.objects.filter(
        user=request.user
    ).order_by('-uploaded_at')

    total_files = files.count()

    storage_used = round(
        sum(file.file_size for file in files),
        2
    )

    total_carbon = round(
        sum(file.carbon_impact for file in files),
        4
    )

    if total_carbon < 1:
        carbon_value = round(
            total_carbon * 1000,
            2
        )
        carbon_unit = 'mg CO₂'
    else:
        carbon_value = round(
            total_carbon,
            2
        )
        carbon_unit = 'g CO₂'

    return render(
        request,
        'dashboard.html',
        {
            'files': files,
            'total_files': total_files,
            'storage_used': storage_used,
            'total_carbon': carbon_value,
            'carbon_unit': carbon_unit
        }
    )
def view_file(request, file_id):

    file = UploadedFile.objects.get(id=file_id)

    if file.user != request.user:
        return redirect('dashboard')

    if file.file_name.lower().endswith('.docx'):

        with file.file.open('rb') as docx_file:
            result = mammoth.convert_to_html(docx_file)

        return render(request, 'view_file.html', {
            'file_name': file.file_name,
            'content': result.value
        })
    if file.file_name.lower().endswith('.pptx'):

        presentation = Presentation(file.file.path)

        slides_content = []

        for slide in presentation.slides:

            slide_text = []

            for shape in slide.shapes:

                if hasattr(shape, "text"):
                    if shape.text.strip():
                        slide_text.append(shape.text)

            slides_content.append(slide_text)

        return render(request, 'view_ppt.html', {
            'file_name': file.file_name,
            'slides': slides_content
        })
    return redirect(file.file.url)
@login_required
def delete_file(request, file_id):

    file = UploadedFile.objects.get(id=file_id)

    if file.user == request.user:

        file.file.delete()
        file.delete()

    return redirect('dashboard')


def logout_view(request):

    logout(request)

    return redirect('login')